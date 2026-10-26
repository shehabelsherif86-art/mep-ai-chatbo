import streamlit as st
from google import genai
from google.genai import types
import tempfile
import os
import time
import hashlib
from io import BytesIO
from docx import Document
from docx.oxml import OxmlElement

# ---------------------------------------------------------
# الإعدادات العامة
# ---------------------------------------------------------
# اسم الموديل (لو جوجل قفلته مستقبلاً غيّره من هنا فقط)
MODEL_NAME = "gemini-3.8-flash"

# موديلات احتياطية لو الأساسي عليه ضغط (خطأ 503). كلها موديلات مختلفة عن الأساسي
# عشان لو واحد عليه ضغط يجرب اللي بعده (Flash-Lite أخف وسعته غالباً أكبر)
FALLBACK_MODELS = ["gemini-3.5-flash-lite", "gemini-3.1-flash-lite"]

# موديلات تحويل الصوت إلى نص (بالترتيب)
TRANSCRIBE_MODELS = [MODEL_NAME, "gemini-3.5-flash-lite", "gemini-3.1-flash-lite"]

st.set_page_config(page_title="المساعد الهندسي الذكي للمواصفات", layout="wide")

# ضبط اتجاه النص (عربي + إنجليزي) وشكل الجداول
st.markdown("""
<style>
.stApp { direction: rtl; }

[data-testid="stMarkdownContainer"] p,
[data-testid="stMarkdownContainer"] li,
[data-testid="stMarkdownContainer"] h1,
[data-testid="stMarkdownContainer"] h2,
[data-testid="stMarkdownContainer"] h3,
[data-testid="stMarkdownContainer"] h4,
[data-testid="stMarkdownContainer"] td,
[data-testid="stMarkdownContainer"] th {
    unicode-bidi: plaintext;
    text-align: start;
}

[data-testid="stMarkdownContainer"] ul,
[data-testid="stMarkdownContainer"] ol { direction: rtl; }
[data-testid="stMarkdownContainer"] table { direction: rtl; width: 100%; }

[data-testid="stMarkdownContainer"] code {
    direction: ltr;
    unicode-bidi: isolate;
    background: transparent;
    color: inherit;
    padding: 0;
}
</style>
""", unsafe_allow_html=True)

SYSTEM_PROMPT = """
أنت استشاري هندسي كهربائي خبير. أجب على أسئلة المهندس بالاعتماد حصرياً على المستندات المرفقة.
يجب أن تكون إجاباتك هندسية دقيقة، ومنظمة في شكل نقاط (Bullet Points)، واستخدم جداول المقارنة عند استعراض الفروق الفنية أو البنود.
اذكر رقم البند أو الصفحة إن أمكن.

قواعد التنسيق (مهمة جداً):
- اكتب بالعربية، والمصطلحات الإنجليزية بين قوسين بعد المصطلح العربي، مثال: كابلات أحادية النواة (Single Core).
- لا تستخدم علامة الدولار ($) ولا صيغ LaTeX نهائياً.
- لا تستخدم علامات الاقتباس العكسية (`) نهائياً.
- اكتب كل الأرقام بالأرقام الإنجليزية (0-9).
- اكتب المقاسات والوحدات كنص عادي، مثال: 4 × 240 مم²، واستخدم جدولاً أو نقاطاً قصيرة بدل جملة عربية طويلة فيها أرقام متتالية.
"""


# ---------------------------------------------------------
# دوال مساعدة
# ---------------------------------------------------------
def clean_text(text):
    """علامة $ بتخلي ستريمليت يعرض الأرقام كمعادلات، والـ backticks بتبوّظ ترتيب العربي."""
    return (text or "").replace("$", "\\$").replace("`", "")


def wait_for_file(client, uploaded):
    """استنى لحد ما الملف المرفوع يبقى جاهز عند جوجل."""
    while uploaded.state.name == "PROCESSING":
        time.sleep(2)
        uploaded = client.files.get(name=uploaded.name)
    if uploaded.state.name == "FAILED":
        raise RuntimeError("فشلت معالجة أحد الملفات عند جوجل. جرّب ملف PDF تاني.")
    return uploaded


def upload_files_to_gemini(client, files):
    refs = []
    for file in files:
        path = None
        try:
            with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as f:
                f.write(file.getvalue())
                path = f.name
            uploaded = client.files.upload(file=path)
            refs.append(wait_for_file(client, uploaded))
        finally:
            if path and os.path.exists(path):
                os.unlink(path)
    return refs


def start_chat(client, file_refs):
    """بيبدأ محادثة جديدة والملفات مرفقة في أول رسالة."""
    file_parts = [
        types.Part.from_uri(file_uri=ref.uri, mime_type=ref.mime_type or "application/pdf")
        for ref in file_refs
    ]
    history = [
        types.Content(
            role="user",
            parts=file_parts + [types.Part.from_text(
                text="هذه هي المستندات الهندسية المرجعية. اعتمد عليها حصرياً في إجاباتك."
            )],
        ),
        types.Content(
            role="model",
            parts=[types.Part.from_text(
                text="لقد اطلعت على المستندات الهندسية المرفقة وأنا جاهز للإجابة على الاستفسارات الفنية."
            )],
        ),
    ]
    return client.chats.create(
        model=MODEL_NAME,
        config=types.GenerateContentConfig(system_instruction=SYSTEM_PROMPT),
        history=history,
    )


def is_temporary_error(error):
    """أخطاء مؤقتة من جوجل (ضغط على الموديل أو حد الاستخدام)."""
    msg = str(error)
    return any(k in msg for k in (
        "503", "UNAVAILABLE", "429", "RESOURCE_EXHAUSTED", "500", "INTERNAL", "overloaded"
    ))


def send_with_retry(client, query):
    """بيعيد المحاولة كذا مرة لو الموديل عليه ضغط، وبعدها بيجرب موديل احتياطي."""
    st.session_state.fallback_used = None
    for attempt in range(3):
        try:
            return st.session_state.chat.send_message(query).text
        except Exception as e:
            if not is_temporary_error(e):
                raise
            time.sleep(2 ** (attempt + 1))  # 2 ثم 4 ثم 8 ثواني

    # كل المحاولات فشلت: جرّب الموديلات الاحتياطية بنفس سجل المحادثة
    config = types.GenerateContentConfig(system_instruction=SYSTEM_PROMPT)
    history = st.session_state.chat.get_history()
    last_error = None
    for model in FALLBACK_MODELS:
        try:
            fallback_chat = client.chats.create(model=model, config=config, history=history)
            answer = fallback_chat.send_message(query).text
            # نرجع للموديل الأساسي في الأسئلة الجاية مع الاحتفاظ بسجل المحادثة
            st.session_state.chat = client.chats.create(
                model=MODEL_NAME, config=config, history=fallback_chat.get_history()
            )
            st.session_state.fallback_used = model
            return answer
        except Exception as e:
            if not is_temporary_error(e):
                raise
            last_error = e
    raise last_error


def transcribe_audio(client, audio_bytes, mime_type="audio/wav"):
    """بيحوّل التسجيل الصوتي إلى نص عربي مكتوب بالحروف العربية."""
    prompt = (
        "هذا تسجيل صوتي لمهندس كهرباء مصري يسأل سؤالاً بالعربية (اللهجة المصرية). "
        "فرّغ الكلام إلى نص مكتوب بالحروف العربية كما قيل بالضبط. "
        "لا تترجم ولا تفسّر ولا تغيّر الكلمات ولا تضف شيئاً. "
        "اكتب الكلمات العربية بحروف عربية حتى لو بدت لك قريبة من كلمات إنجليزية. "
        "لا تكتب كلمة بالإنجليزية إلا إذا نطقها المتحدث فعلاً بالإنجليزية بوضوح "
        "كاختصار أو مصطلح هندسي (مثل MCCB أو XLPE). "
        "إذا لم يكن هناك كلام واضح اكتب فقط: [غير واضح]. "
        "اكتب نص الكلام فقط."
    )
    contents = [types.Part.from_bytes(data=audio_bytes, mime_type=mime_type), prompt]
    config = types.GenerateContentConfig(temperature=0.0)
    last_error = None
    for model in TRANSCRIBE_MODELS:
        try:
            result = client.models.generate_content(model=model, contents=contents, config=config)
            return (result.text or "").strip()
        except Exception as e:
            if not is_temporary_error(e):
                raise
            last_error = e
    raise last_error


def build_word_file(text):
    """بيعمل ملف Word من آخر إجابة، والفقرات من اليمين لليسار."""
    doc = Document()
    doc.add_heading("تقرير استشاري هندسي - مواصفات الكهرباء", 0)
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("#"):
            para = doc.add_heading(line.lstrip("#").strip().replace("**", ""), level=2)
        else:
            para = doc.add_paragraph(line.replace("**", ""))
        para._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# ---------------------------------------------------------
# الواجهة
# ---------------------------------------------------------
st.sidebar.title("⚙️ إعدادات النظام الهندسي")
api_key = st.sidebar.text_input("أدخل مفتاح Google AI Studio API", type="password")

st.title("⚡ مساعد مواصفات وعقود الكهرباء (MEP AI Consultant)")

uploaded_files = st.sidebar.file_uploader(
    "ارفع ملفات المواصفات أو الأكواد (PDF)", type=["pdf"], accept_multiple_files=True
)

if "messages" not in st.session_state:
    st.session_state.messages = []

if uploaded_files and api_key:
    api_key = api_key.strip()
    client = genai.Client(api_key=api_key)

    # لو الملفات أو المفتاح اتغيروا، نرفع الملفات ونبدأ محادثة جديدة
    signature = (tuple((f.name, f.size) for f in uploaded_files), api_key)
    if st.session_state.get("signature") != signature:
        try:
            with st.spinner("جاري قراءة وتحليل المستندات الهندسية..."):
                refs = upload_files_to_gemini(client, uploaded_files)
                st.session_state.chat = start_chat(client, refs)
            st.session_state.signature = signature
            st.session_state.messages = []
        except Exception as e:
            st.session_state.pop("signature", None)
            st.error(f"حدث خطأ أثناء رفع المستندات: {e}")
            st.stop()

    # عرض المحادثة السابقة
    for msg in st.session_state.messages:
        with st.chat_message(msg["role"]):
            st.markdown(clean_text(msg["text"]))

    # الإدخال الصوتي (ميكروفون)
    st.sidebar.markdown("---")
    st.sidebar.subheader("🎤 اسأل بالصوت")
    if "audio_n" not in st.session_state:
        st.session_state.audio_n = 0
    audio = st.sidebar.audio_input(
        "اضغط على الميكروفون، اتكلم، وبعدين اضغط إيقاف",
        key=f"audio_{st.session_state.audio_n}",
    )
    if audio is not None:
        audio_bytes = audio.getvalue()
        audio_id = hashlib.md5(audio_bytes).hexdigest()
        # نعالج كل تسجيل مرة واحدة بس (ستريمليت بيعيد تشغيل الصفحة مع كل تفاعل)
        if st.session_state.get("last_audio_id") != audio_id:
            st.session_state.last_audio_id = audio_id
            try:
                with st.spinner("جاري تحويل صوتك إلى نص..."):
                    heard = transcribe_audio(client, audio_bytes, audio.type or "audio/wav")
                if heard and "غير واضح" not in heard:
                    st.session_state.pending_voice = heard
                    st.session_state.pending_voice_box = heard
                else:
                    st.sidebar.warning("ماسمعتش كلام واضح في التسجيل، جرّب تاني.")
            except Exception as e:
                if is_temporary_error(e):
                    st.sidebar.error("خوادم جوجل عليها ضغط دلوقتي. استنى دقيقة وسجّل تاني.")
                else:
                    st.sidebar.error(f"حدث خطأ أثناء تحويل الصوت: {e}")

    # مراجعة النص المفرّغ قبل الإرسال (عشان تصحح أي كلمة اتفهمت غلط)
    voice_query = None
    if st.session_state.get("pending_voice"):
        st.sidebar.caption("✏️ ده اللي فهمته من صوتك. عدّله لو فيه غلط، وبعدين اضغط إرسال:")
        edited = st.sidebar.text_area(
            "نص سؤالك", key="pending_voice_box", label_visibility="collapsed"
        )
        col_send, col_cancel = st.sidebar.columns(2)
        if col_send.button("إرسال ✅", use_container_width=True):
            voice_query = edited.strip() or None
            st.session_state.pending_voice = None
            st.session_state.audio_n += 1  # بيصفّر الميكروفون للتسجيل الجاي
        if col_cancel.button("إلغاء ❌", use_container_width=True):
            st.session_state.pending_voice = None
            st.session_state.audio_n += 1
            st.rerun()

    # واجهة المحادثة
    text_query = st.chat_input("اكتب استفسارك الفني أو اطلب جدول مقارنة للمواصفات...")
    user_query = text_query or voice_query
    if user_query:
        st.session_state.messages.append({"role": "user", "text": user_query})
        with st.chat_message("user"):
            st.markdown(user_query)

        with st.chat_message("assistant"):
            try:
                with st.spinner("جاري التحليل... (لو الموديل عليه ضغط هعيد المحاولة تلقائياً)"):
                    answer = send_with_retry(client, user_query)
                answer = answer or "لم أتمكن من توليد إجابة، جرّب إعادة صياغة السؤال."
                st.markdown(clean_text(answer))
                if st.session_state.get("fallback_used"):
                    st.caption("ℹ️ الموديل الأساسي كان عليه ضغط، فالإجابة دي من موديل أخف. راجع الأرقام المهمة بنفسك.")
                st.session_state.messages.append({"role": "assistant", "text": answer})
                if voice_query:
                    st.rerun()
            except Exception as e:
                # مانخزنش الخطأ كإجابة، ونشيل السؤال عشان تقدر تسأله تاني
                st.session_state.messages.pop()
                if is_temporary_error(e):
                    st.error("خوادم جوجل عليها ضغط عالي دلوقتي. استنى دقيقة واسأل السؤال تاني.")
                else:
                    st.error(f"حدث خطأ أثناء المعالجة: {e}")

    # زر تصدير آخر استشارة إلى Word
    st.sidebar.markdown("---")
    st.sidebar.subheader("📄 تصدير المخرجات")
    assistant_msgs = [m for m in st.session_state.messages if m["role"] == "assistant"]
    if assistant_msgs:
        st.sidebar.download_button(
            label="📥 تحميل آخر استشارة كملف Word",
            data=build_word_file(assistant_msgs[-1]["text"]),
            file_name="MEP_Technical_Report.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    else:
        st.sidebar.caption("اسأل سؤالاً أولاً، وبعدها هيظهر زر التحميل.")
else:
    st.info("الرجاء إدخال مفتاح الـ API ورفع مستندات الـ PDF من القائمة الجانبية لبدء الاستشارات الفنية.")
